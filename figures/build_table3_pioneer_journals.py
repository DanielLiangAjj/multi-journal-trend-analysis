"""Table 3 — Pioneer journals (top 10 per domain).

Renders as both:
  - PNG: side-by-side methodology + health table for direct figure insertion
  - DOCX: native-Word table for paste-into-manuscript

Pioneer-journal count = number of K=100 topics where a given journal published
the first article in our 2011-2025 corpus.
"""
import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).parent))
from _mirror import mirror_to_main

import pandas as pd
import matplotlib.pyplot as plt
import matplotlib.patches as mpatches
import numpy as np

ROOT = Path(__file__).resolve().parent.parent
OUT_PNG = ROOT / "figures" / "table3_pioneer_journals.png"
OUT_DOCX = Path("/Users/danielliang/Downloads/table3_pioneer_journals.docx")

# Compact journal labels for the figure
JOURNAL_ABBR = {
    "Journal of Biomedical Informatics (JBI)": "JBI",
    "Journal of Medical Systems": "Journal of Medical Systems",
    "Methods of Information in Medicine": "Methods of Information in Medicine",
    "Journal of the American Medical Informatics Association (JAMIA)": "JAMIA",
    "Briefings in Bioinformatics": "Briefings in Bioinformatics",
    "Artificial Intelligence in Medicine": "Artificial Intelligence in Medicine",
    "Nature Methods": "Nature Methods",
    "Computers in Biology and Medicine": "Computers in Biology and Medicine",
    "International Journal of Medical Informatics (IJMI)": "IJMI",
    "IEEE Journal of Biomedical and Health Informatics (J-BHI)": "IEEE J-BHI",
    "Nature Medicine": "Nature Medicine",
    "Database: The Journal of Biological Databases and Curation": "Database",
    "Applied Clinical Informatics": "Applied Clinical Informatics",
    "Journal of Medical Internet Research (JMIR)": "JMIR",
}


def load(domain):
    em = pd.read_csv(ROOT / f"data/visualizations_k100/pi_q1_{domain}_first_emergence.csv")
    counts = em["first_journal"].value_counts().head(10)
    total = len(em)
    return [(JOURNAL_ABBR.get(j, j), int(c), 100*c/total) for j, c in counts.items()], total


def build_png(rows_m, total_m, rows_h, total_h):
    fig, (ax_m, ax_h) = plt.subplots(1, 2, figsize=(14, 6),
                                     gridspec_kw={"wspace": 0.3})

    def draw_table(ax, rows, total, title, colour):
        ax.axis("off")
        ax.set_title(title, fontsize=12, fontweight="bold", pad=10)
        n = len(rows)
        # Header row
        ax.text(0.05, 1.0, "Rank", fontsize=10, fontweight="bold", color="#333", transform=ax.transAxes)
        ax.text(0.20, 1.0, "Journal", fontsize=10, fontweight="bold", color="#333", transform=ax.transAxes)
        ax.text(0.80, 1.0, "n topics", fontsize=10, fontweight="bold", color="#333",
                transform=ax.transAxes, ha="right")
        ax.text(0.95, 1.0, "% of all", fontsize=10, fontweight="bold", color="#333",
                transform=ax.transAxes, ha="right")
        # Header underline
        ax.plot([0.02, 0.98], [0.97, 0.97], color="black", lw=0.8, transform=ax.transAxes)

        max_count = max(c for _, c, _ in rows)
        for i, (jname, cnt, pct) in enumerate(rows):
            y = 0.92 - i * 0.085
            # Rank
            ax.text(0.05, y, f"{i+1}.", fontsize=10, color="#666", transform=ax.transAxes)
            # Journal name
            ax.text(0.20, y, jname, fontsize=10, color="#222", transform=ax.transAxes)
            # Bar (proportional to count)
            bar_w = (cnt / max_count) * 0.30
            ax.add_patch(mpatches.Rectangle(
                (0.49, y - 0.022), bar_w, 0.044,
                color=colour, alpha=0.55, transform=ax.transAxes,
                clip_on=False,
            ))
            # Count + percentage
            ax.text(0.80, y, str(cnt), fontsize=10, fontweight="bold",
                    transform=ax.transAxes, ha="right", color="#222")
            ax.text(0.95, y, f"{pct:.1f}%", fontsize=9.5, color="#666",
                    transform=ax.transAxes, ha="right")
        # Footer: total
        ax.text(0.5, 0.92 - n * 0.085 - 0.04,
                f"Total topics in domain: {total}",
                fontsize=9, color="#666", style="italic",
                transform=ax.transAxes, ha="center")

    draw_table(ax_m, rows_m, total_m,
               "(A) Methodology pioneer journals (top 10 of 70 topics)",
               "#1f4ea0")
    draw_table(ax_h, rows_h, total_h,
               "(B) Health pioneer journals (top 10 of 86 topics)",
               "#9c27b0")

    fig.suptitle(
        "Table 3. Pioneer journals — top 10 per domain (K=100)\n"
        "Pioneer-journal count = number of topics where the journal published the "
        "first article in our 2011–2025 corpus.",
        fontsize=12.5, fontweight="bold", y=0.99,
    )
    fig.text(0.5, 0.02,
        "Caveat: \"first article\" is sensitive to the corpus start year (2011); methodology mean first-year = 2011.8 (max 2018), "
        "health mean first-year = 2011.2 (max 2014). The ranking captures who-published-first within our 2011-2025 window, not absolute origination.",
        fontsize=8.5, color="#666", style="italic", ha="center", wrap=True,
    )
    fig.tight_layout(rect=[0, 0.04, 1, 0.92])
    fig.savefig(OUT_PNG, dpi=200, bbox_inches="tight")
    plt.close(fig)
    print(f"  wrote {OUT_PNG}")
    mirror_to_main(OUT_PNG)


def build_docx(rows_m, total_m, rows_h, total_h):
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    h = doc.add_heading("Table 3. Pioneer journals — top 10 per domain (K=100)", level=2)
    for r in h.runs: r.font.color.rgb = RGBColor(0x14, 0x3a, 0x6b)

    p = doc.add_paragraph()
    r = p.add_run(
        "Pioneer-journal count = number of topics where the journal published the first article "
        "in our 2011-2025 corpus."
    )
    r.italic = True; r.font.size = Pt(10)

    # Side-by-side: 2 cols of 4 sub-cols each, separator in middle
    tbl = doc.add_table(rows=1 + 10, cols=8)
    tbl.style = "Light Grid Accent 1"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["#", "Methodology pioneer", "n", "% of 70", "#", "Health pioneer", "n", "% of 86"]
    for i, h in enumerate(headers):
        c = tbl.rows[0].cells[i]
        c.text = h
        for p in c.paragraphs:
            for r in p.runs: r.bold = True; r.font.size = Pt(10)

    for i in range(10):
        row = tbl.rows[i + 1].cells
        m = rows_m[i] if i < len(rows_m) else (None, "", 0, 0)
        h = rows_h[i] if i < len(rows_h) else (None, "", 0, 0)
        row[0].text = f"{i+1}"
        row[1].text = m[0]
        row[2].text = str(m[1])
        row[3].text = f"{m[2]:.1f}%"
        row[4].text = f"{i+1}"
        row[5].text = h[0]
        row[6].text = str(h[1])
        row[7].text = f"{h[2]:.1f}%"
        for c in row:
            for p in c.paragraphs:
                for r in p.runs: r.font.size = Pt(9.5)

    # Caveat
    p = doc.add_paragraph()
    r = p.add_run(
        "Caveat: \"first article\" is sensitive to the corpus start year (2011); "
        "methodology mean first-year = 2011.8 (max 2018), health mean first-year = 2011.2 "
        "(max 2014). The ranking captures who-published-first within our 2011-2025 window, "
        "not absolute origination."
    )
    r.italic = True; r.font.size = Pt(9.5); r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    OUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_DOCX)
    print(f"  wrote {OUT_DOCX}")


def main():
    rows_m, total_m = load("methodology")
    rows_h, total_h = load("health")
    print("Top 10 methodology pioneers:")
    for r in rows_m: print(f"  {r}")
    print("Top 10 health pioneers:")
    for r in rows_h: print(f"  {r}")
    print()
    build_png(rows_m, total_m, rows_h, total_h)
    build_docx(rows_m, total_m, rows_h, total_h)


if __name__ == "__main__":
    main()
